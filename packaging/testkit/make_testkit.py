#!/usr/bin/env python3
"""VNCCad: build the engineers' test kit documents from scenarios.json.

Writes HUONG_DAN_TEST.html (how to install, run and report) and PHIEU_TEST.xlsx (the test
sheet to fill in). Usage: python3 make_testkit.py VERSION [OUT_DIR]
"""
import html
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
VERSION = sys.argv[1] if len(sys.argv) > 1 else "dev"
OUT = sys.argv[2] if len(sys.argv) > 2 else HERE
WEB = "https://maitrieuquang-max.github.io/VNCCAD/"
REPO = "https://github.com/maitrieuquang-max/VNCCAD"

rows = json.load(open(os.path.join(HERE, "scenarios.json"), encoding="utf-8"))
groups = []
for r in rows:
    if r[0] not in groups:
        groups.append(r[0])

# ------------------------------------------------------------------ Excel test sheet
wb = Workbook()
F = "Arial"
thin = Side(style="thin", color="B7B7B7")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F3A5F")
input_fill = PatternFill("solid", fgColor="FFF2CC")
wrap = Alignment(wrap_text=True, vertical="top")

ws0 = wb.active
ws0.title = "Hướng dẫn"
info = [
    ("PHIẾU TEST VNCCad", None),
    ("Phiên bản", VERSION),
    ("Người test", ""),
    ("Đơn vị / bộ phận", ""),
    ("Ngày test", ""),
    ("Máy (Windows/Linux, RAM, card đồ họa)", ""),
    ("Bản dùng (Desktop / Web, trình duyệt)", ""),
    (None, None),
    ("Cách điền", "Sheet \"Kịch bản\": chỉ điền các ô nền vàng — cột Kết quả (chọn Đạt / Lỗi / Chưa test / Không áp dụng), Mô tả lỗi, Tệp/ảnh kèm. Dòng 1 là ví dụ."),
    ("Báo lỗi", "Mỗi lỗi ghi: các bước làm, kết quả thấy được, kết quả mong đợi; kèm ảnh chụp màn hình và tệp bản vẽ nếu được. Lỗi nặng (mất dữ liệu, treo, mở sai bản vẽ) ghi rõ \"NẶNG\"."),
    ("Gửi lại", "Gửi phiếu này cùng ảnh/tệp kèm cho người phụ trách VNCCad."),
    ("Tải bản mới", REPO + "/releases"),
    ("Bản web", WEB),
]
for i, (k, v) in enumerate(info, start=1):
    if k is None:
        continue
    a = ws0.cell(row=i, column=1, value=k)
    a.font = Font(name=F, bold=True, size=14 if i == 1 else 11)
    if v is not None:
        b = ws0.cell(row=i, column=2, value=v)
        b.font = Font(name=F)
        b.alignment = wrap
        if 3 <= i <= 7:
            b.fill = input_fill
            b.border = box
ws0.column_dimensions["A"].width = 38
ws0.column_dimensions["B"].width = 100

ws = wb.create_sheet("Kịch bản")
heads = ["STT", "Nhóm", "Chức năng", "Các bước", "Kết quả mong đợi", "Bản", "Kết quả", "Mô tả lỗi / ghi chú", "Tệp / ảnh kèm"]
widths = [6, 13, 26, 60, 42, 10, 13, 45, 22]
for c, (h, w) in enumerate(zip(heads, widths), start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(name=F, bold=True, color="FFFFFF")
    cell.fill = head_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = box
    ws.column_dimensions[cell.column_letter].width = w
dv = DataValidation(type="list", formula1='"Đạt,Lỗi,Chưa test,Không áp dụng"', allow_blank=True)
ws.add_data_validation(dv)
for i, r in enumerate(rows, start=2):
    vals = [i - 1, r[0], r[1], r[2], r[3], r[4], "", "", ""]
    if i == 2:
        vals[6], vals[7] = "Đạt", "Ví dụ: mở trong 3 giây"
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.font = Font(name=F, size=10)
        cell.alignment = wrap
        cell.border = box
        if c >= 7:
            cell.fill = input_fill
    dv.add(f"G{i}")
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:I{len(rows) + 1}"
last = len(rows) + 1

ws2 = wb.create_sheet("Tổng hợp")
for c, h in enumerate(["Nhóm", "Số kịch bản", "Đạt", "Lỗi", "Chưa test", "Tỷ lệ đạt"], start=1):
    cell = ws2.cell(row=1, column=c, value=h)
    cell.font = Font(name=F, bold=True, color="FFFFFF")
    cell.fill = head_fill
    cell.border = box
for i, g in enumerate(groups, start=2):
    rng_g = f"'Kịch bản'!$B$2:$B${last}"
    rng_k = f"'Kịch bản'!$G$2:$G${last}"
    vals = [
        g,
        f"=COUNTIF({rng_g},A{i})",
        f'=COUNTIFS({rng_g},A{i},{rng_k},"Đạt")',
        f'=COUNTIFS({rng_g},A{i},{rng_k},"Lỗi")',
        f"=B{i}-C{i}-D{i}-COUNTIFS({rng_g},A{i},{rng_k},\"Không áp dụng\")",
        f"=IF(C{i}+D{i}=0,\"-\",C{i}/(C{i}+D{i}))",
    ]
    for c, v in enumerate(vals, start=1):
        cell = ws2.cell(row=i, column=c, value=v)
        cell.font = Font(name=F)
        cell.border = box
    ws2.cell(row=i, column=6).number_format = "0%"
t = len(groups) + 2
ws2.cell(row=t, column=1, value="Tổng").font = Font(name=F, bold=True)
for c, col in enumerate("BCDE", start=2):
    cell = ws2.cell(row=t, column=c, value=f"=SUM({col}2:{col}{t - 1})")
    cell.font = Font(name=F, bold=True)
cell = ws2.cell(row=t, column=6, value=f"=IF(C{t}+D{t}=0,\"-\",C{t}/(C{t}+D{t}))")
cell.font = Font(name=F, bold=True)
cell.number_format = "0%"
ws2.cell(row=t + 2, column=1, value="Tỷ lệ đạt = Đạt / (Đạt + Lỗi); các kịch bản chưa test không tính.").font = Font(name=F, italic=True, size=9)
ws2.column_dimensions["A"].width = 16
for col in "BCDEF":
    ws2.column_dimensions[col].width = 13
xlsx = os.path.join(OUT, "PHIEU_TEST.xlsx")
wb.save(xlsx)

# ------------------------------------------------------------------ HTML guide
def esc(s):
    return html.escape(s)

items = []
for g in groups:
    lis = "".join(
        f"<tr><td>{esc(r[1])}</td><td>{esc(r[2])}</td><td>{esc(r[3])}</td><td class=b>{esc(r[4])}</td></tr>" for r in rows if r[0] == g
    )
    items.append(f"<h3>{esc(g)}</h3><table><thead><tr><th>Chức năng</th><th>Các bước</th><th>Kết quả mong đợi</th><th>Bản</th></tr></thead><tbody>{lis}</tbody></table>")

page = f"""<!doctype html>
<html lang="vi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Hướng dẫn test VNCCad {esc(VERSION)}</title>
<style>
:root {{ --ink:#1d2733; --dim:#5b6776; --line:#d9dee5; --accent:#1f5fa8; --bg:#fbfcfd; --pill:#eef3f9; }}
* {{ box-sizing:border-box }}
body {{ margin:0; background:var(--bg); color:var(--ink); font:15px/1.55 "Segoe UI", Arial, sans-serif; }}
main {{ max-width:980px; margin:0 auto; padding:28px 18px 60px; }}
h1 {{ font-size:26px; margin:0 0 4px }} .ver {{ color:var(--dim); margin:0 0 22px }}
h2 {{ font-size:19px; margin:30px 0 10px; padding-bottom:6px; border-bottom:2px solid var(--line) }}
h3 {{ font-size:16px; margin:20px 0 8px; color:var(--accent) }}
.grid {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(400px,1fr)); gap:12px }}
.card {{ background:#fff; border:1px solid var(--line); border-radius:8px; padding:14px 16px }}
.card b {{ display:block; margin-bottom:4px }}
code, kbd {{ background:var(--pill); border-radius:4px; padding:1px 5px; font-family:Consolas, monospace; font-size:13.5px }}
table {{ width:100%; border-collapse:collapse; background:#fff; font-size:13.5px; margin-bottom:6px }}
th, td {{ border:1px solid var(--line); padding:6px 8px; vertical-align:top; text-align:left }}
th {{ background:var(--pill) }} td.b {{ white-space:nowrap; color:var(--dim) }}
.note {{ border-left:4px solid #e0a100; background:#fff8e6; padding:10px 14px; border-radius:4px }}
ol li, ul li {{ margin:3px 0 }}
@media print {{ body {{ background:#fff }} main {{ padding:0 }} .card {{ break-inside:avoid }} }}
</style></head><body><main>
<h1>Hướng dẫn test VNCCad</h1>
<p class="ver">Phiên bản {esc(VERSION)} · phần mềm CAD 2D kiểu AutoCAD, chạy trên máy (offline) và trên trình duyệt (online)</p>

<h2>1. Lấy phần mềm</h2>
<div class="grid">
<div class="card"><b>Windows (offline)</b>Giải nén <code>VNCCad-{esc(VERSION)}-windows-x64.zip</code> vào một thư mục, chạy <code>VNCCad.exe</code>. Không cần cài đặt.<br>Lần đầu Windows có thể hiện <i>“Windows protected your PC”</i> (phần mềm chưa ký số): bấm <b style="display:inline">More info → Run anyway</b>.</div>
<div class="card"><b>Linux (offline)</b>Giải nén <code>VNCCad-{esc(VERSION)}-linux-x86_64.tar.gz</code>, chạy <code>./vnccad</code>. Cần môi trường đồ họa (X11 hoặc Wayland).</div>
<div class="card"><b>Trình duyệt (online)</b>Mở <a href="{WEB}">{WEB}</a> bằng Chrome hoặc Edge. Bấm biểu tượng <i>Cài đặt</i> trên thanh địa chỉ để dùng như ứng dụng; sau lần mở đầu, chạy được cả khi mất mạng.</div>
<div class="card"><b>Máy chủ nội bộ (tùy chọn)</b><code>VNCCad-{esc(VERSION)}-web.zip</code> là trang web tĩnh: chép lên máy chủ web nội bộ (cần HTTPS để chạy offline). Không mở trực tiếp <code>index.html</code> bằng cách bấm đúp.</div>
</div>

<h2>2. Bộ test kèm theo (thư mục <code>mau</code>)</h2>
<ul>
<li><code>ban_ve_mau.dwg</code> / <code>.dxf</code> — dầm I có hatch và kích thước, tường MLINE, mảng cọc D600, bảng tọa độ, tim tuyến; Layout1/Layout2 có khung tên thuộc tính (thử Bộ bản vẽ).</li>
<li><code>nhap_xuat_toado.lsp</code> — lisp mẫu: <code>NHAPTD</code> (nhập cọc từ CSV), <code>XUATTD</code> (xuất tọa độ ra CSV), <code>DAYCS</code> (dây cao su <code>grread</code>), <code>LOPMAU</code> (tạo lớp bằng ActiveX).</li>
<li><code>coc_khoan_nhoi.csv</code> — tọa độ 6 cọc cho <code>NHAPTD</code>.</li>
<li><code>PHIEU_TEST.xlsx</code> — phiếu ghi kết quả.</li>
</ul>
<p>Nạp lisp/CSV: kéo-thả tệp vào cửa sổ VNCCad (hoặc lệnh <code>APPLOAD</code>).</p>

<h2>3. Cách test và báo lỗi</h2>
<ol>
<li>Điền thông tin người test và máy ở sheet <i>Hướng dẫn</i> của <code>PHIEU_TEST.xlsx</code>.</li>
<li>Làm lần lượt các kịch bản ở mục 4 (cùng nội dung với sheet <i>Kịch bản</i>); chọn kết quả <b style="display:inline">Đạt / Lỗi / Chưa test</b>.</li>
<li>Ưu tiên test với <b style="display:inline">bản vẽ thật</b> đang dùng ở dự án, so sánh với AutoCAD.</li>
<li>Mỗi lỗi ghi: các bước làm, thấy gì, mong đợi gì; kèm ảnh màn hình và tệp bản vẽ nếu được. Lỗi mất dữ liệu, treo máy, lưu ra AutoCAD không mở được ghi rõ <b style="display:inline">NẶNG</b>.</li>
<li>Sheet <i>Tổng hợp</i> tự tính tỷ lệ đạt theo nhóm.</li>
</ol>
<p class="note">Đây là bản thử nghiệm. Với bản vẽ quan trọng, luôn giữ bản gốc và làm việc trên bản sao.</p>

<h2>4. Kịch bản test</h2>
{''.join(items)}

<h2>5. Phím và lệnh hay dùng</h2>
<p><kbd>Esc</kbd> hủy lệnh · <kbd>Enter</kbd>/<kbd>Space</kbd> lặp lệnh · <kbd>Ctrl+Z</kbd> hoàn tác · <kbd>F3</kbd> bắt điểm · <kbd>F8</kbd> ortho · <kbd>Ctrl+1</kbd> bật/tắt bảng bên phải · <kbd>Ctrl+2</kbd> DesignCenter · bí danh lệnh như AutoCAD (<code>L</code>, <code>PL</code>, <code>C</code>, <code>TR</code>, <code>O</code>…). Danh sách đầy đủ: tệp <code>VNCCAD.md</code> trong gói.</p>
</main></body></html>
"""
with open(os.path.join(OUT, "HUONG_DAN_TEST.html"), "w", encoding="utf-8") as f:
    f.write(page)
print("wrote", xlsx, "and HUONG_DAN_TEST.html")
